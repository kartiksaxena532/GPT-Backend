from flask import Flask, request, send_file
from flask_cors import CORS  # Import CORS
from docx import Document
from utils.docx_editor import replace_placeholders
from utils.patent_scraper import scrape_google_patent
from utils.generate_concepts import generate_overview_and_concepts
from utils.flask_integration import edit_template_action  # Import the action from flask_integration.py

import os
from datetime import datetime

app = Flask(__name__)

# Enable CORS for all routes and all origins
CORS(app)  # This will allow all domains to access your API

# Folder paths for templates and output
TEMPLATES_FOLDER = os.path.join(os.getcwd(), 'templates')
OUTPUT_FOLDER = os.path.join(os.getcwd(), 'output')

# Ensure output directory exists
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

@app.route('/edit-template', methods=['POST'])
def edit_template():
    # Get data from JSON body
    data = request.get_json()

    # Ensure required fields are present in the request
    required_fields = ['template_name', 'name', 'patent', 'cutoff_date']
    for field in required_fields:
        if field not in data:
            return f'Missing required field: {field}', 400

    # Extract data from the request
    template_name = data['template_name']
    name = data['name']
    patent_id = data['patent']
    cutoff_date = data['cutoff_date']

    # Find the template file in the templates folder
    template_path = os.path.join(TEMPLATES_FOLDER, f"{template_name}.docx")
    if not os.path.exists(template_path):
        return f"Template file '{template_name}.docx' not found in templates folder.", 404

    try:
        # Open the template file
        doc = Document(template_path)
    except Exception as e:
        return f'Error opening template file: {str(e)}', 400

    # Base replacements
    replacements = {
        'name': name,
        'date': datetime.today().strftime('%d-%m-%Y'),
        'company': 'GreyB',  # Default, can be replaced if needed
        'patent': patent_id,
        'cutoff_date': cutoff_date,
        'priority_date': 'N/A'
    }

    # Auto-generate overview, in-depth concept, and priority date if patent ID is provided
    if patent_id and patent_id != 'ExamplePatent':
        try:
            patent_data = scrape_google_patent(patent_id)
            concepts = generate_overview_and_concepts(
                patent_data['description'],
                patent_data['claims']
            )
            replacements['overview'] = concepts.get('overview', 'N/A')
            replacements['in_depth_concept'] = concepts.get('in_depth_concept', 'N/A')
            replacements['priority_date'] = patent_data.get('priority_date', 'N/A')
        except Exception as e:
            print("Patent scraping or concept generation failed:", e)
            replacements['overview'] = 'Failed to generate overview.'
            replacements['in_depth_concept'] = 'Failed to generate in-depth concept.'
            replacements['priority_date'] = 'Failed to retrieve priority date.'

    # Replace placeholders in doc
    replace_placeholders(doc, replacements)

    # Save the updated document to the output folder
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_filename = f"updated_{template_name}_{timestamp}.docx"
    output_path = os.path.join(OUTPUT_FOLDER, output_filename)
    doc.save(output_path)

    # Return the modified document as a response
    return send_file(output_path, as_attachment=True, download_name=output_filename)

if __name__ == '__main__':
    app.run(debug=False, host='0.0.0.0', port=5000)
