import re
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl

def replace_placeholders(doc, replacements):
    # Body content
    for para in doc.paragraphs:
        _replace_runs_safe(para.runs, replacements)

    for table in doc.tables:
        _replace_table_safe(table, replacements)

    # Headers and footers (use XML to fully access structure)
    for section in doc.sections:
        _replace_in_hf_element(section.header._element, doc, replacements)
        _replace_in_hf_element(section.footer._element, doc, replacements)

    # Shapes, text boxes, etc.
    _replace_in_shapes(doc, replacements)

def _replace_table_safe(table, replacements):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                _replace_runs_safe(para.runs, replacements)
            for nested in cell.tables:
                _replace_table_safe(nested, replacements)

def _replace_runs_safe(runs, replacements):
    """
    Replaces placeholders that might be split across multiple runs.
    Preserves formatting from the first run.
    """
    i = 0
    while i < len(runs):
        if '{{' in runs[i].text:
            combined = runs[i].text
            indices = [i]
            j = i + 1

            while '}}' not in combined and j < len(runs):
                combined += runs[j].text
                indices.append(j)
                j += 1

            match = re.search(r'{{(.*?)}}', combined)
            if match:
                placeholder = match.group(0)
                key = match.group(1).strip()
                value = replacements.get(key, placeholder)

                # Preserve formatting
                fmt = runs[indices[0]]
                formatting = {
                    "bold": fmt.bold,
                    "italic": fmt.italic,
                    "underline": fmt.underline,
                    "font": fmt.font.name,
                    "size": fmt.font.size
                }

                # Clear placeholder text
                for idx in indices:
                    runs[idx].text = ''

                # Write replacement
                runs[indices[0]].text = value
                runs[indices[0]].bold = formatting["bold"]
                runs[indices[0]].italic = formatting["italic"]
                runs[indices[0]].underline = formatting["underline"]
                runs[indices[0]].font.name = formatting["font"]
                runs[indices[0]].font.size = formatting["size"]

                i = indices[-1] + 1
            else:
                i += 1
        else:
            i += 1

def _replace_in_shapes(doc, replacements):
    """
    Finds and replaces placeholders inside shapes and text boxes using raw XML.
    """
    for shape in doc.element.body.iter():
        if isinstance(shape, CT_P):
            para = Paragraph(shape, doc)
            _replace_runs_safe(para.runs, replacements)
        elif isinstance(shape, CT_Tbl):
            for row in shape.xpath(".//w:tr"):
                for cell in row.xpath(".//w:tc"):
                    for para_el in cell.xpath(".//w:p"):
                        para = Paragraph(para_el, doc)
                        _replace_runs_safe(para.runs, replacements)

def _replace_in_hf_element(hf_element, doc, replacements):
    """
    Parses header/footer XML structure to replace placeholders.
    """
    for elem in hf_element.iter():
        if isinstance(elem, CT_P):
            para = Paragraph(elem, doc)
            _replace_runs_safe(para.runs, replacements)
        elif isinstance(elem, CT_Tbl):
            for row in elem.xpath(".//w:tr"):
                for cell in row.xpath(".//w:tc"):
                    for para_el in cell.xpath(".//w:p"):
                        para = Paragraph(para_el, doc)
                        _replace_runs_safe(para.runs, replacements)
